from pathlib import Path

from selenium.common import NoSuchElementException, UnexpectedAlertPresentException, TimeoutException
from selenium.webdriver.support import expected_conditions as EC
import time
from PyQt5.QtCore import QTimer
from selenium.webdriver.common.by import By
from selenium.webdriver.support.wait import WebDriverWait
from openpyxl import Workbook, load_workbook
from config.paths import EXCEL_PROCESSOS_PATH
from config.settings import WAIT_MEDIUM, URL_DESTINO


def tentar_clicar_botao_pdf(
    self,
    driver,
    worker_id,
    numero
):

    try:

        self.log(
            f"📄 Worker {worker_id} procurando PDF"
        )

        elemento = WebDriverWait(driver, 20).until(
            EC.element_to_be_clickable(
                (
                    By.XPATH,
                    "//div[@id='389' or @id='394']/ancestor::tr//img[contains(@class,'salvaDocumento')]"
                )
            )
        )



        if not elemento:

            self.log(
                f"❌ Worker {worker_id} PDF não encontrado"
            )

            return None

        elemento.click()

        self.log(
            f"🖱 Worker {worker_id} clicou PDF"
        )

        WebDriverWait(driver, WAIT_MEDIUM).until(
            EC.presence_of_element_located(
                (
                    By.ID,
                    "janelaModalCaptchaDownload"
                )
            )
        )

        self.log(
            f"🧩 Worker {worker_id} captcha aberto"
        )

        return self.tratar_modal_captcha(
            driver,
            worker_id
        )

    except Exception as e:

        self.log(
            f"❌ Worker {worker_id} erro PDF: {e}"
        )

        return None

def selecionar_processo(self, item):
    if self.processando:
        return  # ignora clique durante processamento em lote
    numero = item.text()
    self.entry_processo.setText(numero)
    self.log(f"📌 Processo selecionado: {numero}")
def _registrar_processo_concluido(self, numero):
    if not numero:
        self.log_new("⚠️ Tentativa de registrar processo concluído com número inválido.")
        return
    with open(self.arquivo_concluidos, "a", encoding="utf-8") as f:
        f.write(str(numero) + "\n")
    self.processos_concluidos.add(numero)

def atualizar_lista_processos(self):
    self.lista_processos.clear()
    self.processos = []
    try:
        if not EXCEL_PROCESSOS_PATH.exists():
            raise FileNotFoundError(
                f"Arquivo não encontrado: {EXCEL_PROCESSOS_PATH}"
            )
        wb = load_workbook(EXCEL_PROCESSOS_PATH, data_only=True)
        ws = wb.active
        header = ws["A1"].value
        if header != "Numero_Processo":
            raise ValueError(
                "A coluna A deve se chamar 'Numero_Processo'"
            )
        vistos = set()
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if not row or not row[0]:
                continue
            numero = str(row[0]).strip()
            # ignora duplicados e concluídos
            if numero in vistos or numero in self.processos_concluidos:
                continue
            vistos.add(numero)
            self.processos.append(numero)
            self.lista_processos.addItem(numero)
        self.total_processos = len(self.processos)
        self._atualizar_contador_ui()
        self.log(f"📄 {len(self.processos)} processos carregados do Excel.")
    except Exception as e:
        self.ui_error(
            "Erro Excel",
            f"Falha ao carregar processos do Excel"
        )
        self.log_new(f"Falha ao carregar processos do Excel")

def possui_servico_389_ou_394(self, driver) -> bool:
    """
    Verifica se existe serviço 389 ou 394 na tabela de PDFs
    """
    try:
        driver.find_element(
            By.XPATH,
            "//a[normalize-space()='389' or normalize-space()='394']"
        )
        self.log("📄 Serviço 389 ou 394 detectado")
        return True
    except NoSuchElementException:
        self.log_new("📄 Serviço 389/394 NÃO encontrado")
        return False

def _repetir_processo_atual(self, worker_id, motivo):

    numero = self.numero_atual[worker_id]

    self.log_new(
        f"🔁 REPEAT worker {worker_id} "
        f"processo {numero}"
    )

    self.log_new(f"📌 Motivo retry: {motivo}")

    self.tentativas_por_processo.setdefault(numero, 0)

    self.tentativas_por_processo[numero] += 1

    self.log_new(
        f"📌 Tentativa atual: "
        f"{self.tentativas_por_processo[numero]}"
    )

    if self.tentativas_por_processo[numero] >= self.MAX_TENTATIVAS:

        self.log_new(
            f"❌ Processo {numero} excedeu tentativas"
        )

        self.numero_atual[worker_id] = None

        self.processando[worker_id] = False

        self.log_new(
            f"🔄 Chamando próximo worker {worker_id}"
        )

        QTimer.singleShot(
            0,
            lambda: self._processar_proximo(worker_id)
        )

        return

    self.numero_atual[worker_id] = None

    self.processando[worker_id] = False

    self.log_new(
        f"⏳ Worker {worker_id} aguardará 3s antes retry"
    )

    QTimer.singleShot(
        3000,
        lambda: self._processar_proximo(worker_id)
    )

def abrir_detalhe_processo(self, driver, worker_id):
    try:

        numero = self.numero_atual.get(worker_id)

        if not numero:
            self.log_new(f"❌ Worker {worker_id} sem número")
            return

        # ✅ Aborta imediatamente se já foi concluído (evita reprocessar após retry)
        if numero in self.processos_concluidos:
            self.log(f"⏭️ Worker {worker_id} processo {numero} já concluído — pulando")
            self._finalizar_processo_atual(worker_id)
            return

        self.log_new(
            f"🚀 PIPELINE START worker {worker_id} | {numero}"
        )

        # ==========================================
        # LOGIN
        # ==========================================
        self.garantir_login(driver)

        # ==========================================
        # PESQUISA
        # ==========================================
        driver.get(URL_DESTINO)

        campo = WebDriverWait(driver, 30).until(
            EC.visibility_of_element_located(
                (By.NAME, "NumPedido")
            )
        )

        campo.clear()
        campo.send_keys(numero)
        campo.submit()

        # detalhe
        try:
            link = WebDriverWait(driver, 20).until(
                EC.element_to_be_clickable(
                    (
                        By.CSS_SELECTOR,
                        "a[href*='detail']"
                    )
                )
            )
        except TimeoutException:
            # 🔧 FIX: quando o INPI nao encontra o processo pesquisado, ele
            # mostra uma pagina de "Nenhum resultado" sem link de detalhe
            # nenhum — o WebDriverWait acima sempre estourava os 20s
            # inteiros esperando por um link que nunca ia aparecer, e
            # depois ainda gastava as tentativas de retry (MAX_TENTATIVAS)
            # antes de desistir, sem nunca marcar o processo como
            # concluido (ou seja, ele seria tentado de novo na proxima
            # execucao). Detectando esse texto aqui, tratamos como
            # concluido de uma vez, na hora.
            if "Nenhum resultado foi encontrado" in driver.page_source:
                self.log(
                    f"⏭️ Worker {worker_id} processo {numero} "
                    f"sem resultado no INPI — marcando concluído"
                )
                self._registrar_processo_concluido(numero)
                self._finalizar_processo_atual(worker_id)
                return
            raise

        driver.execute_script(
            "arguments[0].click();",
            link
        )

        WebDriverWait(driver, 20).until(
            EC.presence_of_element_located(
                (By.TAG_NAME, "body")
            )
        )

        time.sleep(2)

        # serviço
        if not self.possui_servico_389_ou_394(driver):
            self.log(
                f"⏭️ Worker {worker_id} sem serviço"
            )

            self._registrar_processo_concluido(numero)

            self._finalizar_processo_atual(worker_id)

            return
        # ==========================================
        # PETIÇÕES
        # ==========================================
        self.log(
            f"📂 Worker {worker_id} liberando petições"
        )

        selenium = self._selenium_por_worker(worker_id)

        selenium.bloquear_fechamento_abas = True

        try:
            # 🔧 FIX: liberar_acesso_peticiones so deve rodar quando
            # garantir_acesso_peticiones realmente clicou no link de
            # restricao (ou seja, uma popup vai aparecer). Antes ela
            # rodava sempre, e quando o processo ja tinha acesso liberado
            # (sem popup nenhuma) ela ficava esperando os 20s completos
            # do WAIT_POPUP a toa e, desde o fix anterior que propaga a
            # falha como erro, isso derrubava o processo inteiro e forcava
            # um retry desnecessario — travando o fluxo normal ("so baixa
            # se passar pelo popup").
            precisa_popup = self.garantir_acesso_peticiones(driver)
            if precisa_popup:
                self.liberar_acesso_peticiones(driver)
            else:
                self.log(f"⏭️ Worker {worker_id} sem restrição de petições — pulando popup")
        finally:
            selenium.bloquear_fechamento_abas = False

        self.log(
            f"✅ Worker {worker_id} petições liberadas"
        )

        time.sleep(2)
        # pdf
        caminho_pdf = self._fluxo_pdf(
            driver,
            worker_id,
            numero
        )

        if not caminho_pdf:
            raise Exception("PDF não baixado")

        # renomear
        self._renomear_pdf_para_processo(
            worker_id,
            caminho_pdf
        )

        # finalizar
        self._finalizar_processo_atual(worker_id)

    except Exception as e:

        import traceback

        self.log_new(traceback.format_exc())

        # 🔧 FIX: se a sessão do navegador morreu (crash do Chrome, sessão
        # inválida, driver desconectado), o retry sozinho nunca resolvia —
        # ficava tentando contra o mesmo driver morto para sempre. Agora
        # detectamos esses casos e reiniciamos o navegador antes do retry.
        erro_str = str(e).lower()
        sinais_sessao_morta = (
            "invalid session id",
            "session deleted",
            "session not created",
            "chrome not reachable",
            "disconnected: not connected to devtools",
            "target window already closed",
            "no such window",
            "connection refused",
            "chrome crashed",
            "read timed out",
            "timed out",
            "connection aborted",
            "remote end closed connection",
            "max retries exceeded",
        )
        if any(sinal in erro_str for sinal in sinais_sessao_morta):
            self.log_new(
                f"💀 Worker {worker_id}: sessão do navegador morta detectada — reiniciando navegador"
            )
            try:
                self.reiniciar_sessao_worker(worker_id)
            except Exception as e2:
                self.log_new(
                    f"❌ Worker {worker_id}: falha ao reiniciar navegador: {e2}"
                )

        self._repetir_processo_atual(
            worker_id,
            str(e)
        )
def _fluxo_pdf(self, driver, worker_id, numero):

    self.log(
        f"📥 Worker {worker_id} iniciando fluxo PDF"
    )

    # ✅ Verifica se o PDF já foi baixado antes de repetir todo o fluxo
    nome_seguro = "".join(c for c in str(numero) if c.isalnum())
    pasta_worker = Path(self.download_dirs[worker_id])
    pdf_existente = pasta_worker / f"{nome_seguro}.pdf"
    if pdf_existente.exists():
        self.log(f"✅ Worker {worker_id} PDF já existe: {pdf_existente.name} — pulando download")
        return str(pdf_existente)

    caminho_pdf = self.tentar_clicar_botao_pdf(
        driver,
        worker_id,
        numero
    )

    self.log(
        f"📥 Worker {worker_id} retorno captcha: {caminho_pdf}"
    )

    if not caminho_pdf:

        self.log(
            f"❌ Worker {worker_id} PDF não encontrado"
        )

        return None

    caminho_pdf = Path(caminho_pdf)

    self.log(
        f"📄 Worker {worker_id} verificando arquivo: {caminho_pdf}"
    )

    if not caminho_pdf.exists():

        self.log(
            f"❌ Worker {worker_id} arquivo inexistente"
        )

        return None

    self.log(
        f"✅ Worker {worker_id} PDF OK"
    )

    return caminho_pdf